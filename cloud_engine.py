"""Generate the GM visibility workbook on Linux without Excel or COM."""
from copy import copy,deepcopy
from pathlib import Path
import re
import math
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.cell_range import CellRange,MultiCellRange
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.views import Selection
from openpyxl.workbook.properties import CalcProperties
from report_engine import INPUTS,OUTPUTS,LAYOUTS,align_rows,source_totals,col,month_add
from buy_plan import PLAN_SHEETS,PLAN_COLUMNS,plan_headers,normalize_plan,match_key,report_populations,header_for_column
from metadata import SCHEMAS,MAPS,KEY_LENGTH,normalize_metadata
from output_sort import SORT_FIELDS,index
from cloud_excel import read_sheet,Calculator,save_with_cache
from number_formats import format_report_numbers

WIDTHS={'GRC_DATA':34,'PPO_DATA':40,'GM_TNA':26,'MNL_PPO':39}

def capture_row(sheet,row,width):
    return [(sheet.cell(row,c).value,copy(sheet.cell(row,c)._style),copy(sheet.cell(row,c).alignment)) for c in range(1,width+1)]

def translated(value,old,new):
    # Subtotal ranges are resized explicitly after rebuilding the detail rows;
    # translating an old footer upward can otherwise create negative references.
    if isinstance(value,str) and 'SUBTOTAL(' in value.upper():return value
    return Translator(value,origin=old).translate_formula(new) if isinstance(value,str) and value.startswith('=') else value

def set_row(sheet,row,source,origin,values=None):
    for c,(value,style,alignment) in enumerate(source,1):
        cell=sheet.cell(row,c);cell._style=copy(style);cell.alignment=copy(alignment)
        cell.value=values[c-1] if values is not None else translated(value,f'{col(c)}{origin}',cell.coordinate)

def resize_cf(sheet,oldlast,last):
    existing=[(copy(cf),[deepcopy(rule) for rule in sheet.conditional_formatting[cf]]) for cf in sheet.conditional_formatting]
    sheet.conditional_formatting=ConditionalFormattingList()
    for cf,rules in existing:
        areas=[]
        for area in cf.sqref.ranges:
            area=copy(area)
            if area.max_row==oldlast:area.max_row=max(last,area.min_row)
            if area.max_row==oldlast+1:area.max_row=max(last+1,area.min_row)
            if area.min_row==5 and last<5:continue
            areas.append(str(area))
        if areas:
            for rule in rules:sheet.conditional_formatting.add(' '.join(areas),rule)

def configure_output_headers(sheet,name,month):
    if name.endswith('DEPT'):sheet['M3' if name.startswith('REGULAR') else 'J3']='ATTRIBUTE'
    removable={'REGULAR_ART':[('AU','ATTRIBUTE')],'WINTER_DEPT':[('AC',"STH% JA'26"),('AB',"STH% MAMJ'26")],'WINTER_ART':[('AE',"STH% JA'26"),('AD',"STH% MAMJ'26"),('AC','ATTRIBUTE')]}
    for column,label in removable.get(name,[]):
        if str(sheet[column+'3'].value or '').strip().upper()==label.upper():sheet.delete_cols(index(column)+1)
    for row in sheet.iter_rows(min_row=1,max_row=3):
        for cell in row:
            if cell.data_type!='f' and isinstance(cell.value,str):cell.value=header_for_column(name,cell.column_letter,cell.value,month)

def report_rows(sheet,name,population,labels,plans,totals,masters):
    l=LAYOUTS[name];width=sheet.max_column;oldlast=4
    while sheet.cell(oldlast+1,index(l['keys'][0])+1).value not in (None,''):oldlast+=1
    originals={match_key([sheet[c+str(r)].value for c in l['keys']]):(r,capture_row(sheet,r,width)) for r in range(5,oldlast+1)}
    merchant_column='I' if name.startswith('REGULAR') else 'F'
    merchants={k[:3]:source[index(merchant_column)][0] for k,(_,source) in originals.items()}
    prototype=capture_row(sheet,5,width);footer=capture_row(sheet,oldlast+1,width)
    height=sheet.row_dimensions[5].height;prepared=[]
    for category in population:
        original=originals.get(category);origin,source=original if original else (5,prototype)
        values=[v if original or (isinstance(v,str) and v.startswith('=')) else None for v,_,_ in source]
        for c,value in zip(l['keys'],labels.get(category,list(category))):values[index(c)]=value
        if not original:values[index(merchant_column)]=merchants.get(category[:3])
        if name.endswith('DEPT') and str(sheet['M3' if name.startswith('REGULAR') else 'J3'].value or '').upper()=='PREFERENCE':values[index('M' if name.startswith('REGULAR') else 'J')]=None
        for c,metric in ((l['grc'],'grc'),(l['tna'],'tna'),(l['manual'],'manual'),(l['ppo'],'regular_ppo' if l['months'] else 'winter_ppo')):values[index(c)]=totals[(metric,category)]
        for i,c in enumerate(l['months']):values[index(c)]=totals[(f'ppo_{i}',category)]
        plan=plans.get(category,[0.0]*(3 if name.startswith('REGULAR') else 2))
        for i,c in enumerate(PLAN_COLUMNS[name]):values[index(c)]=sum(plan) if name.startswith('REGULAR') and i==3 else plan[i]
        if masters is not None:
            for dest,(master,field) in MAPS[name].items():
                record=masters[master].get(category[:KEY_LENGTH[master]])
                values[index(dest)]=record[SCHEMAS[master].index(field)] if record else None
        prepared.append((category,origin,source,values))
    def sortkey(item):
        values=item[3]
        return tuple(-float(values[index(c)] or 0) if order==2 else str(values[index(c)] or '').strip().casefold() for c,order in SORT_FIELDS[name])
    prepared.sort(key=sortkey)
    # Remove previous body/footer cells and their row dimensions, retaining headers,
    # column widths, outline definitions, page setup and style patterns.
    for coord in list(sheet._cells):
        if coord[0]>=5:del sheet._cells[coord]
    for r in list(sheet.row_dimensions):
        if r>=5:del sheet.row_dimensions[r]
    for r,(_,origin,source,values) in enumerate(prepared,5):
        values=[translated(v,f'{col(c)}{origin}',f'{col(c)}{r}') for c,v in enumerate(values,1)]
        set_row(sheet,r,source,origin,values);sheet.row_dimensions[r].height=height
    last=4+len(prepared);end=last+1
    set_row(sheet,end,footer,oldlast+1)
    for row in (2,end):
        for cell in sheet[row]:
            if cell.data_type=='f' and 'SUBTOTAL(' in cell.value.upper():cell.value=f'=SUBTOTAL(9,{cell.column_letter}5:{cell.column_letter}{last})' if last>=5 else '=0'
    resize_cf(sheet,oldlast,last);sheet.print_area=f'A1:{col(width)}{end}'
    merchant=merchant_column+'5';pane=sheet.sheet_view.pane
    active='bottomLeft' if pane and (pane.xSplit or 0)>=index(merchant_column)+1 and pane.ySplit else 'bottomRight' if pane and pane.ySplit else None
    sheet.sheet_view.selection=[Selection(pane=active,activeCell=merchant,sqref=merchant)]
    if pane and active:pane.activePane=active
    sheet.sheet_state='visible'
    return last

def refresh_report(template,uploads,destination,report_month,as_of,tna_field='TNA_Q ',progress=lambda s:None,buy_plans=None,metadata_path=None,summary_destination=None):
    if as_of<report_month:raise ValueError('GRC cutoff must be on or after the starting month.')
    if not buy_plans or set(buy_plans)!= {'REGULAR','WINTER'}:raise ValueError('Upload both Regular and Winter buy plans.')
    if Path(template).suffix.lower()!='.xlsx':raise ValueError('Cloud report templates must be XLSX. Save your XLSB reference as XLSX in Excel first.')
    book=openpyxl.load_workbook(template);input_rows={};counts={};warnings=[];plan_values={};plan_rows={};plan_labels={}
    original_keys={n:{match_key([book[n][c+str(r)].value for c in LAYOUTS[n]['keys']]) for r in range(5,book[n].max_row) if book[n][LAYOUTS[n]['keys'][0]+str(r)].value not in (None,'')} for n in OUTPUTS}
    for name in INPUTS:
        progress('Reading '+name)
        canonical=[book[name].cell(1,c).value for c in range(1,WIDTHS[name]+1)]
        if name=='PPO_DATA':canonical[39]='DELIVERY DATE'
        path,selected=uploads[name];rows,ignored=align_rows(read_sheet(path,selected,True),canonical,name,tna_field)
        quantity={'GRC_DATA':31,'PPO_DATA':38,'GM_TNA':23 if '®' in tna_field else 22,'MNL_PPO':37}[name]
        if any(not math.isfinite(float(row[quantity])) for row in rows):raise ValueError(f'{name}: quantities must be finite numeric values.')
        input_rows[name]=rows;counts[name]=len(rows)
        if ignored:warnings.append(f'{name}: ignored {ignored} rows without category keys.')
        old=book[name];book.remove(old);sheet=book.create_sheet(name);sheet.append(canonical)
        for row in rows:sheet.append(row)
        sheet.sheet_state='hidden'
    for group in ('REGULAR','WINTER'):
        for name in OUTPUTS:
            if not name.startswith(group):continue
            rows,values,_=normalize_plan(read_sheet(buy_plans[group],PLAN_SHEETS[name]),name,report_month,[])
            plan_values[name]=values;plan_rows[PLAN_SHEETS[name]]=len(rows)
            for row in rows:plan_labels[match_key(row[:len(LAYOUTS[name]['keys'])])]=row[:len(LAYOUTS[name]['keys'])]
            label=PLAN_SHEETS[name]
            if label in book:book.remove(book[label])
            sheet=book.create_sheet(label);sheet.append(plan_headers(name,report_month))
            for row in rows:sheet.append(row)
            sheet.sheet_state='hidden'
    masters=None;metadata_counts={}
    if metadata_path:
        masters={}
        for label in SCHEMAS:
            rows,_=normalize_metadata(read_sheet(metadata_path,label),label);metadata_counts[label]=len(rows)
            masters[label]={match_key(row[:KEY_LENGTH[label]]):row for row in rows}
            if label in book:book.remove(book[label])
            sheet=book.create_sheet(label);sheet.append(SCHEMAS[label]+['LOOKUP_KEY'])
            for row in rows:sheet.append(row[:-1]+[' | '.join(match_key(row[:KEY_LENGTH[label]]))])
            sheet.sheet_state='hidden'
    populations,labels=report_populations(input_rows,plan_values,original_keys);labels.update(plan_labels)
    totals=source_totals(input_rows,report_month,as_of,'X' if '®' in tna_field else 'W');lastrows={};caches={};records={}
    for name in OUTPUTS:
        progress('Generating '+name)
        # Preference is cleared before the Attribute label is applied.
        lastrows[name]=report_rows(book[name],name,populations[name],labels,plan_values[name],totals,masters)
        configure_output_headers(book[name],name,report_month)
        format_report_numbers(book[name],lastrows[name])
        calculator=Calculator(book[name]);caches[name]=calculator.calculate_all()
        records[name]=[[calculator.cell(f'{col(c)}{r}') for c in range(1,book[name].max_column+1)] for r in range(5,lastrows[name]+1)]
    for sheet in book:
        if sheet.title not in OUTPUTS:sheet.sheet_state='hidden'
    book.active=book.sheetnames.index(OUTPUTS[0]);book.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True)
    progress('Saving checked XLSX report with live formulas and calculated values')
    save_with_cache(book,destination,caches)
    summary_stats={}
    if summary_destination:
        try:
            from cloud_summary import create_summary
            progress('Generating value-pasted Regular / Winter / Bottom summary')
            summary_stats=create_summary(Path(__file__).parent/'template/Summary_Template.xlsx',summary_destination,records,report_month,as_of)
        except Exception as error:
            Path(summary_destination).unlink(missing_ok=True);summary_stats={'error':str(error)};warnings.append('Summary unavailable: '+str(error))
    return {'rows':counts,'buy_plan_rows':plan_rows,'metadata_rows':metadata_counts,'summary':summary_stats,'output_rows':{n:lastrows[n]-4 for n in OUTPUTS},'warnings':warnings,'differences':[],'formula_errors':[],'report_month':str(report_month),'as_of':str(as_of),'validation':'Cloud calculation completed; imported fields and totals are values, report business formulas remain live; complete summary sheets are values.'}
