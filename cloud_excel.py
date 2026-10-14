"""Portable Excel input readers and safe calculation/cache support for report formulas."""
import ast
import math
import re
import zipfile
from functools import lru_cache
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl

class FormulaError(ValueError):pass

def read_sheet(path,name,single_fallback=False):
    path=Path(path);suffix=path.suffix.lower()
    if suffix=='.xlsb':
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as book:
            selected=name if name in book.sheets else book.sheets[0] if single_fallback and len(book.sheets)==1 else None
            if selected is None:raise ValueError(f'{path.name}: worksheet {name} is missing.')
            with book.get_sheet(selected) as sheet:
                rows=[]
                for cells in sheet.rows():
                    if not cells:continue
                    while len(rows)<=cells[0].r:rows.append([])
                    row=[None]*(max(c.c for c in cells)+1)
                    for c in cells:row[c.c]=c.v
                    rows[cells[0].r]=row
                return rows
    if suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(str(path))
        if book.datemode:raise ValueError('1904-date workbooks are unsupported. Save using the 1900 date system.')
        selected=name if name in book.sheet_names() else book.sheet_names()[0] if single_fallback and book.nsheets==1 else None
        if selected is None:raise ValueError(f'{path.name}: worksheet {name} is missing.')
        sheet=book.sheet_by_name(selected)
        return [[xlrd.xldate_as_datetime(c.value,book.datemode) if c.ctype==xlrd.XL_CELL_DATE else c.value for c in sheet.row(r)] for r in range(sheet.nrows)]
    if suffix not in ('.xlsx','.xlsm'):raise ValueError('Upload XLSX, XLSB or XLS files.')
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        if book.epoch.year==1904:raise ValueError('1904-date workbooks are unsupported. Save using the 1900 date system.')
        selected=name if name in book.sheetnames else book.sheetnames[0] if single_fallback and len(book.sheetnames)==1 else None
        if selected is None:raise ValueError(f'{path.name}: worksheet {name} is missing.')
        return [list(row) for row in book[selected].iter_rows(values_only=True)]
    finally:book.close()

@lru_cache(maxsize=30000)
def parse_formula(formula):
    if len(formula)>10000:raise FormulaError('Formula is too long.')
    parts=re.split(r'("(?:[^"]|"")*")',formula.lstrip('='))
    for i in range(0,len(parts),2):
        text=parts[i].replace('$','')
        text=re.sub(r'\b([A-Z]{1,3}\d+):([A-Z]{1,3}\d+)\b',lambda m:f'RANGE("{m[1]}","{m[2]}")',text)
        # Preserve reference strings generated for RANGE while replacing cells.
        fragments=re.split(r'("[^"]*")',text)
        for j in range(0,len(fragments),2):fragments[j]=re.sub(r'\b([A-Z]{1,3}\d+)\b',lambda m:f'CELL("{m[1]}")',fragments[j])
        text=''.join(fragments).replace('<>','!=').replace('^','**')
        text=re.sub(r'(?<![<>=!])=(?!=)','==',text)
        parts[i]=text
    try:return ast.parse(''.join(parts),mode='eval').body
    except SyntaxError as error:raise FormulaError(f'Unsupported report formula: {formula}') from error

class Calculator:
    def __init__(self,sheet):self.sheet=sheet;self.cache={};self.active=set()
    @staticmethod
    def numeric(value):
        if value is None or value=='':return 0
        if isinstance(value,(int,float)):return value
        raise ValueError(f'Non-numeric report input {value!r}.')
    def cell(self,address):
        if address in self.cache:return self.cache[address]
        value=self.sheet[address].value
        if not (isinstance(value,str) and value.startswith('=')):return value
        if address in self.active:raise FormulaError(f'Circular report formula at {address}.')
        self.active.add(address)
        try:
            result=self.evaluate(parse_formula(value))
            if isinstance(result,(int,float)) and not math.isfinite(result):raise FormulaError('Non-finite result.')
            self.cache[address]=result;return result
        except FormulaError:raise
        except Exception as error:raise FormulaError(f'{self.sheet.title}!{address}: {error}') from error
        finally:self.active.remove(address)
    def evaluate(self,node):
        if isinstance(node,ast.Constant):return node.value
        if isinstance(node,ast.UnaryOp):
            n=self.numeric(self.evaluate(node.operand))
            if isinstance(node.op,ast.USub):return -n
            if isinstance(node.op,ast.UAdd):return n
        if isinstance(node,ast.BinOp):
            a=self.numeric(self.evaluate(node.left));b=self.numeric(self.evaluate(node.right))
            if isinstance(node.op,ast.Add):return a+b
            if isinstance(node.op,ast.Sub):return a-b
            if isinstance(node.op,ast.Mult):return a*b
            if isinstance(node.op,ast.Div):return a/b
            if isinstance(node.op,ast.Pow):return a**b
        if isinstance(node,ast.Compare) and len(node.ops)==1:
            a=self.numeric(self.evaluate(node.left));b=self.numeric(self.evaluate(node.comparators[0]));op=node.ops[0]
            if isinstance(op,ast.Gt):return a>b
            if isinstance(op,ast.Lt):return a<b
            if isinstance(op,ast.GtE):return a>=b
            if isinstance(op,ast.LtE):return a<=b
            if isinstance(op,ast.Eq):return a==b
            if isinstance(op,ast.NotEq):return a!=b
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and not node.keywords:
            name=node.func.id.upper();args=node.args
            if name=='CELL':return self.cell(self.evaluate(args[0]))
            if name=='RANGE':
                return [self.cell(c.coordinate) for row in self.sheet[f'{self.evaluate(args[0])}:{self.evaluate(args[1])}'] for c in row]
            if name=='IF':return self.evaluate(args[1] if self.evaluate(args[0]) else args[2])
            if name=='IFERROR':
                try:return self.evaluate(args[0])
                except FormulaError:raise
                except (ZeroDivisionError,ValueError,TypeError,OverflowError):return self.evaluate(args[1])
            if name in ('MIN','MAX','SUM','SUBTOTAL'):
                values=[self.evaluate(a) for a in args]
                if name=='SUBTOTAL':
                    if values[0] not in (9,109):raise FormulaError('Only SUM subtotals are supported.')
                    values=values[1:]
                flat=[self.numeric(x) for v in values for x in (v if isinstance(v,list) else [v])]
                if name in ('SUM','SUBTOTAL'):return sum(flat)
                return (min(flat) if name=='MIN' else max(flat)) if flat else 0
            raise FormulaError(f'Unsupported report function {name}. Use the provided XLSX template.')
        raise FormulaError('Unsupported report formula syntax. Use the provided XLSX template.')
    def calculate_all(self):
        for row in self.sheet:
            for cell in row:
                if cell.data_type=='f':self.cell(cell.coordinate)
        return self.cache

def save_with_cache(book,path,caches):
    """Keep live formulas plus cloud-computed cached results for Excel/viewers."""
    book.save(path);ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    temporary=Path(path).with_suffix('.cached.xlsx')
    with zipfile.ZipFile(path) as source,zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as target:
        sheets={f'xl/worksheets/sheet{i}.xml':caches.get(sheet.title,{}) for i,sheet in enumerate(book.worksheets,1)}
        for info in source.infolist():
            data=source.read(info.filename)
            if info.filename in sheets and sheets[info.filename]:
                tree=ET.fromstring(data)
                for cell in tree.findall(f'.//{{{ns}}}c'):
                    address=cell.get('r')
                    if address not in sheets[info.filename] or cell.find(f'{{{ns}}}f') is None:continue
                    value=sheets[info.filename][address]
                    cached=cell.find(f'{{{ns}}}v')
                    if cached is None:cached=ET.SubElement(cell,f'{{{ns}}}v')
                    if isinstance(value,bool):cell.set('t','b');cached.text='1' if value else '0'
                    elif isinstance(value,(int,float)):cell.attrib.pop('t',None);cached.text=repr(value)
                    else:raise FormulaError(f'Unexpected non-numeric formula result at {address}.')
                data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            target.writestr(info,data)
    temporary.replace(path)
