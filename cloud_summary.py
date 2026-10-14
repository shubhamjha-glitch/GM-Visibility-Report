"""Template-formatted summary snapshots computed from the checked report rows."""
from copy import copy
import openpyxl
from openpyxl.styles import Font
from report_engine import month_add
from summary_engine import SPEC,column_index,group_names,bottom_selection
from number_formats import keep_cell_number_formats

def clone_cell(source,target):
    target.value=None if source.data_type=='f' else source.value
    copy_style(source,target)
    font=copy(source.font);font.sz=8;target.font=font

def copy_style(source,target):
    for field in ('font','fill','border','alignment','protection'):setattr(target,field,copy(getattr(source,field)))
    target.number_format=source.number_format

def copy_layout(original,target,rows):
    target.sheet_format=copy(original.sheet_format);target.sheet_properties=copy(original.sheet_properties)
    target.page_setup=copy(original.page_setup);target.page_margins=copy(original.page_margins)
    target.print_options=copy(original.print_options);target.freeze_panes=original.freeze_panes
    for c,dimension in original.column_dimensions.items():
        target.column_dimensions[c]=copy(dimension);target.column_dimensions[c].parent=target
        copy_style(dimension,target.column_dimensions[c])
    mapping={}
    for new,old in enumerate(rows,1):
        mapping.setdefault(old,new)
        for cell in original[old]:clone_cell(cell,target.cell(new,cell.column))
        target.row_dimensions[new]=copy(original.row_dimensions[old]);target.row_dimensions[new].index=new
        target.row_dimensions[new].parent=target
        copy_style(original.row_dimensions[old],target.row_dimensions[new])
    for area in original.merged_cells.ranges:
        if area.min_row in mapping and area.max_row in mapping:
            target.merge_cells(start_row=mapping[area.min_row],end_row=mapping[area.max_row],start_column=area.min_col,end_column=area.max_col)
    target.print_area=f'C4:Q{len(rows)}'
    target.page_setup.orientation='landscape';target.page_setup.fitToWidth=1;target.page_setup.fitToHeight=1
    target.sheet_properties.pageSetUpPr.fitToPage=True

def measures(records,name):
    s=SPEC[name]
    total=lambda column:sum(float(row[column_index(column)] or 0) for row in records)
    denominators=[total(s['bp']),total(s['total']),total(s['bp']),total(s['total']),total(s['total'])]
    ratios=[total(c)/d if d else (1 if i==4 and name.startswith('REGULAR') else 0) for i,(c,d) in enumerate(zip('ABCDE',denominators))]
    return ratios+[total(s['coverage'])]

def fill_summary(original,sheet,group,records,month,asof):
    names=(group+'_DEPT',group+'_ART');buyers=group_names(records,names,'merchant');divisions=group_names(records,names,'division')
    rows=list(range(1,11))+[11]*len(buyers)+list(range(15,20))+[20]*len(divisions)+list(range(28,original.max_row+1))
    copy_layout(original,sheet,rows)
    delta=len(buyers)-4;division_start=20+delta
    sheet['D5']='REGULAR & SUMMER' if group=='REGULAR' else 'WINTER'
    for datecell in ('E8','K8',f'E{17+delta}',f'K{17+delta}'):sheet[datecell]='Till '+asof.strftime('%d-%b-%Y').upper()
    period=''.join(month_add(month,i).strftime('%b')[0].upper() for i in range(3)) if group=='REGULAR' else 'WINTER'
    for header in (10,19+delta):
        for c in ('E','K'):sheet[f'{c}{header}']=month.strftime('%b').upper()+'\nFILL RATE %'
        for c in ('F','L'):sheet[f'{c}{header}']=period+'\nFILL RATE %'
        for c in ('G','M'):sheet[f'{c}{header}']=month.strftime('%b').upper()+'\nVISIBILITY %'
        for c in ('H','N'):sheet[f'{c}{header}']=period+'\nVISIBILITY %'
    for start,labels,field in ((11,buyers,'merchant'),(division_start,divisions,'division')):
        for offset,label in enumerate(labels):
            r=start+offset;sheet.cell(r,4,label or 'UNASSIGNED')
            for name,first in zip(names,(5,11)):
                subset=[row for row in records[name] if str(row[column_index(SPEC[name][field])] or '').strip()==label]
                for j,value in enumerate(measures(subset,name)):sheet.cell(r,first+j,value)
        grand=start+len(labels);sheet.cell(grand,4,'GRAND TOTAL')
        for name,first in zip(names,(5,11)):
            for j,value in enumerate(measures(records[name],name)):sheet.cell(grand,first+j,value)
        for r in range(start,grand+1):
            for c in range(5,17):sheet.cell(r,c).number_format='#,##0' if c in (10,16) else '0%'
    return {'buyers':len(buyers),'divisions':len(divisions)}

def create_summary(template,destination,records,month,asof):
    reference=openpyxl.load_workbook(template);book=openpyxl.Workbook();book.loaded_theme=reference.loaded_theme;book.remove(book.active);stats={}
    for group in ('REGULAR','WINTER'):
        stats[group]=fill_summary(reference['REGULAR'],book.create_sheet(group),group,records,month,asof)
    bottom=book.create_sheet('BOTTOM');original=reference['BOTTOM']
    copy_layout(original,bottom,list(range(1,original.max_row+1)));bottom.print_area='B2:V19'
    for row in bottom.iter_rows(min_row=5,max_row=19,min_col=3,max_col=22):
        for c in row:c.value=None
    bottom['B2']='Bottom 15 Departments — Lowest '+month.strftime('%b %Y')+' Visibility % (Sale Rank 1–100)'
    for c in ('G','I','L','O'):bottom[c+'4']=str(bottom[c+'4'].value or '').replace('OCT',month.strftime('%b').upper())
    period=''.join(month_add(month,i).strftime('%b')[0].upper() for i in range(3))
    for c in ('H','M','P'):bottom[c+'4']=str(bottom[c+'4'].value or '').replace('OND',period)
    mapping={'C':'I','D':'J','E':'L','F':'P','G':'Q','H':'T','I':'U','J':'V','K':'W','L':'X','M':'Y','O':'Z','P':'AA','R':'AC','S':'AG','T':'AM','U':'AQ','V':'AR'}
    selected=bottom_selection(records['REGULAR_DEPT'])
    for r,item in enumerate(selected,5):
        values=item[4]
        for target,source in mapping.items():bottom[f'{target}{r}']=values[column_index(source)]
    for r in range(5,20):
        for c in range(6,23):bottom.cell(r,c).number_format='0%' if c in (12,13,15,16,18) else '#,##0'
    for sheet in book:
        keep_cell_number_formats(sheet)
        for row in sheet:
            for cell in row:
                font=copy(cell.font);font.sz=8;cell.font=font
                if cell.data_type=='f':raise ValueError(f'{sheet.title}!{cell.coordinate}: summary formula was not value-pasted.')
    book.save(destination);stats['bottom_rows']=len(selected)
    return stats
